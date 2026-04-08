"""
Generate docs/AI_scientist_审题人登记表.xlsx by model.

One model => one sheet. Each sheet lists a fixed scenario roster in one table.
Sheet titles are limited to 31 characters.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Missing dependency: openpyxl. Install with: pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)

ROOT = Path(__file__).resolve().parents[1]
META_RUNS = ROOT / "meta_runs"
OUT = ROOT / "docs" / "AI_scientist_审题人登记表.xlsx"

INNER = Path("round_001") / "inner_workspaces" / "round_001"

# Excel sheet title: max 31 chars; cannot contain []:*?/\\
_INVALID_SHEET = re.compile(r'[\[\]:*?/\\]')
_RUN_MODEL = re.compile(r"^new_(.+?)_no-reviewer_\d{8}_\d{6}$")


def safe_sheet_title(name: str) -> str:
    s = _INVALID_SHEET.sub("_", name)
    if len(s) <= 31:
        return s
    return s[:31]


def list_scenario_ids(run_dir: Path) -> list[str]:
    inner = run_dir / INNER
    if not inner.is_dir():
        return []
    ids = [p.name for p in inner.iterdir() if p.is_dir()]
    return sort_scenario_ids_abc_adjacent(ids)


def model_key_from_run_id(run_id: str) -> str:
    """
    Parse run id like:
    new_Pro_moonshotai_Kimi-K2.5_no-reviewer_20260407_213322
    -> Pro_moonshotai_Kimi-K2.5
    """
    m = _RUN_MODEL.match(run_id)
    if not m:
        return run_id
    return m.group(1)


# Leading pattern e.g. 01a_FooBar, 11c_FooBar
_SCENARIO_HEAD = re.compile(r"^(\d+)([a-z])?_")


def sort_scenario_ids_abc_adjacent(ids: list[str]) -> list[str]:
    """
    同一题号下 a、b、c 紧挨：先题号，再变体字母。
    即 01a, 01b, 01c, 02a, 02b, 02c, …（完整 33 题时共 33 行）。
    """

    def key(sid: str) -> tuple:
        m = _SCENARIO_HEAD.match(sid)
        if not m:
            return (10**9, 99, sid)
        n = int(m.group(1))
        letter = m.group(2)
        vi = ord(letter) - ord("a") if letter else 0
        return (n, vi, sid)

    return sorted(ids, key=key)


def variant_letter_from_scenario_id(sid: str) -> str | None:
    m = _SCENARIO_HEAD.match(sid)
    if not m:
        return None
    return m.group(2) or None


def build_sheet(wb: Workbook, model_key: str, scenario_ids: list[str]) -> None:
    title = safe_sheet_title(model_key)
    # avoid duplicate truncated titles
    existing = {ws.title for ws in wb.worksheets}
    base = title
    n = 1
    while title in existing:
        suffix = f"_{n}"
        title = (base[: 31 - len(suffix)] + suffix)[:31]
        n += 1

    ws = wb.create_sheet(title=title)

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    last_col = 7
    last_data_row = 2 + len(scenario_ids)

    ws["A1"] = "模型（聚合）"
    ws["A1"].font = header_font
    ws["B1"] = model_key
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_col)
    ws["B1"].alignment = Alignment(vertical="center", wrap_text=True)

    headers = [
        "序号",
        "scenario_id（场景）",
        "审题人",
        "领题时间",
        "提交审查时间",
        "负责人确认时间",
        "状态",
    ]
    status_hint = "已领 / 已审 / 负责人已确认"

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = header_font
        cell.alignment = center

    # 浅色按 a/b/c 区分行（同一题号下 abc 连续时三色相邻）
    fill_a = PatternFill(start_color="E8F4FC", end_color="E8F4FC", fill_type="solid")
    fill_b = PatternFill(start_color="EEF8E8", end_color="EEF8E8", fill_type="solid")
    fill_c = PatternFill(start_color="FFF5E6", end_color="FFF5E6", fill_type="solid")
    fills = {"a": fill_a, "b": fill_b, "c": fill_c}

    for i, sid in enumerate(scenario_ids, start=1):
        r = 2 + i
        ws.cell(row=r, column=1, value=i).alignment = center
        ws.cell(row=r, column=2, value=sid)
        vl = variant_letter_from_scenario_id(sid)
        row_fill = fills.get(vl) if vl else None
        if row_fill:
            for c in range(1, last_col + 1):
                ws.cell(row=r, column=c).fill = row_fill

    note_row = last_data_row + 1
    ws.cell(row=note_row, column=1, value=f"状态说明：{status_hint}")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=last_col)

    widths = [6, 52, 14, 18, 18, 18, 12]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "C3"


def main() -> None:
    if not META_RUNS.is_dir():
        print(f"No meta_runs at {META_RUNS}", file=sys.stderr)
        raise SystemExit(1)

    run_dirs = sorted([p for p in META_RUNS.iterdir() if p.is_dir()], key=lambda p: p.name)
    if not run_dirs:
        print("meta_runs is empty; creating one template sheet.", file=sys.stderr)
        run_dirs = []

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    if not run_dirs:
        build_sheet(wb, "TEMPLATE_model_placeholder", ["01a_scenario", "01b_scenario", "01c_scenario"])
    else:
        # Build one global scenario roster; each model sheet uses the same rows
        # so the reviewer table is stable at 33 rows (when abc are complete).
        all_scenarios: set[str] = set()
        model_keys: set[str] = set()
        for rd in run_dirs:
            model_keys.add(model_key_from_run_id(rd.name))
            for sid in list_scenario_ids(rd):
                all_scenarios.add(sid)

        scenario_roster = sort_scenario_ids_abc_adjacent(list(all_scenarios))
        if not scenario_roster:
            scenario_roster = ["（未找到场景目录，请检查路径）"]

        for model_key in sorted(model_keys):
            build_sheet(wb, model_key, scenario_roster)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
